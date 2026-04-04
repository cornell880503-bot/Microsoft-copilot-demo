"""
Project Copilot — Python FastAPI Sidecar
=========================================
Runs alongside the Electron app on http://localhost:8765

Endpoints:
  GET  /health            — liveness probe
  GET  /get-active-window — foreground app title
  POST /index-docs        — index local_data/*.txt into ChromaDB
  POST /search-docs       — semantic similarity search
  POST /agent/run         — full agent pipeline (SSE streaming)
"""

import logging
import os
import sys
import tempfile
from contextlib import asynccontextmanager

from pathlib import Path
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel, Field

import json as _json

from agent import run_agent_stream, generate_chat_title, generate_suggestions
from memory_store import MemoryStore
from privacy_mode import load_privacy_mode, save_privacy_mode, VALID_MODES
from email_sender import send_email
from window_context import get_active_window_title
from rag.indexer import index_local_data
from rag.searcher import search_docs
from action_ledger import ledger
import chats as chats_store

load_dotenv(Path(__file__).parent / ".env")

# ── Logging ───────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("copilot.server")


# ── Lifespan ──────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    ledger.clear_backups()   # clear stale backups from previous session
    logger.info("Copilot sidecar starting on http://localhost:8765")
    yield
    logger.info("Copilot sidecar shutting down")


# ── App ───────────────────────────────────────────────────────────────────────

app = FastAPI(
    title="Project Copilot Sidecar",
    version="1.0.0",
    description="System context & RAG layer for the Copilot Electron shell",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "app://.", "file://"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    errors = exc.errors()
    logger.error("Request validation error on %s: %s", request.url.path, errors)
    # Return detail as a string so the frontend can display it
    detail = "; ".join(
        f"{'.'.join(str(l) for l in e['loc'])}: {e['msg']}" for e in errors
    )
    return JSONResponse(status_code=422, content={"detail": detail})


# ── Schemas ───────────────────────────────────────────────────────────────────

class SearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=1000)
    top_k: int = Field(5, ge=1, le=20)

class SearchResult(BaseModel):
    content: str
    source: str
    score: float

class SearchResponse(BaseModel):
    query: str
    results: list[SearchResult]
    total: int

class IndexRequest(BaseModel):
    extra_dirs: list[str] = Field(default_factory=list, description="Extra absolute folder paths to index")

class HistoryMessage(BaseModel):
    role: str       # "user" | "assistant"
    content: str

class AgentRequest(BaseModel):
    query:           str = Field(..., min_length=1, max_length=2000)
    history:         list[HistoryMessage] = Field(default_factory=list)
    conversation_id: str | None = Field(None)

class SaveFileRequest(BaseModel):
    filename: str = Field(..., min_length=1)
    content:  str = Field(...)

class SendEmailRequest(BaseModel):
    to:              str = Field(..., min_length=1)
    subject:         str = Field(...)
    body:            str = Field(...)
    attachment_path: str | None = Field(None)


# ── Routes ────────────────────────────────────────────────────────────────────

@app.get("/health")
async def health():
    return {"status": "ok", "service": "copilot-sidecar"}


@app.get("/get-active-window")
async def get_active_window():
    title = get_active_window_title()
    if title is None:
        raise HTTPException(status_code=503, detail="Could not determine active window")
    return {"active_window": title}


@app.get("/suggest")
async def suggest(window: str = ""):
    """Return 3 proactive action suggestions based on the active window."""
    active = window or get_active_window_title() or ""
    privacy = load_privacy_mode()
    result = await generate_suggestions(active, privacy_mode=privacy["mode"])
    return {"window": active, "privacy_mode": privacy["mode"], "suggestions": result}


class PrivacyModeRequest(BaseModel):
    mode: str = Field(..., min_length=1)


@app.get("/privacy-mode")
async def get_privacy_mode():
    return load_privacy_mode()


@app.put("/privacy-mode")
async def update_privacy_mode(body: PrivacyModeRequest):
    if body.mode.strip().lower() not in VALID_MODES:
        raise HTTPException(status_code=400, detail=f"Invalid privacy mode: {body.mode}")
    return save_privacy_mode(body.mode)


@app.get("/memory")
async def get_memory():
    store = MemoryStore()
    return store.memory


@app.delete("/memory/{key}")
async def delete_memory(key: str):
    store = MemoryStore()
    deleted = store.delete(key)
    if not deleted:
        raise HTTPException(status_code=404, detail=f"Memory key not found: {key}")
    return {"ok": True, "deleted": key}


@app.post("/index-docs")
async def index_docs(body: IndexRequest = IndexRequest()):
    try:
        result = index_local_data(extra_dirs=body.extra_dirs or None)
        logger.info("Indexing complete: %s", result)
        return result
    except Exception as e:
        logger.exception("Indexing failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/search-docs", response_model=SearchResponse)
async def search_documents(body: SearchRequest):
    try:
        results = search_docs(body.query, top_k=body.top_k)
        return SearchResponse(
            query=body.query,
            results=[SearchResult(**r) for r in results],
            total=len(results),
        )
    except Exception as e:
        logger.exception("Search failed")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/send-email")
async def send_email_endpoint(body: SendEmailRequest):
    """Send a real email via Gmail/Outlook SMTP with optional file attachment."""
    logger.info("send-email called: to=%r subject=%r attachment_path=%r", body.to, body.subject, body.attachment_path)
    try:
        result = send_email(
            to=body.to,
            subject=body.subject,
            body=body.body,
            attachment_path=body.attachment_path,
        )
        logger.info("Email sent: %s", result)
        return result
    except Exception as e:
        logger.exception("Failed to send email")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/save-file")
async def save_file(body: SaveFileRequest):
    """Write content to ~/Downloads/<filename>. Returns the saved path."""
    # Sanitize filename — strip path separators to prevent directory traversal
    safe_name = Path(body.filename).name
    if not safe_name:
        raise HTTPException(status_code=400, detail="Invalid filename")

    save_path = Path.home() / "Downloads" / safe_name
    try:
        # Snapshot before writing so the user can undo
        entry = ledger.snapshot(str(save_path), "SAVE_FILE", safe_name)
        ext = save_path.suffix.lower()
        if ext == ".docx":
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            # Use filename stem (without extension) as title
            title_text = save_path.stem
            doc.add_heading(title_text, level=0)
            for line in body.content.splitlines():
                stripped = line.strip()
                if not stripped:
                    doc.add_paragraph("")
                elif stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], level=3)
                elif stripped.startswith("- ") or stripped.startswith("• "):
                    p = doc.add_paragraph(stripped[2:], style="List Bullet")
                else:
                    doc.add_paragraph(stripped)
            doc.save(str(save_path))
        else:
            save_path.write_text(body.content, encoding="utf-8")
        ledger.push(entry)
        logger.info("Saved file: %s", save_path)
        return {"saved_to": str(save_path), "filename": safe_name}
    except Exception as e:
        logger.exception("Failed to save file")
        raise HTTPException(status_code=500, detail=str(e))


class DeleteFileRequest(BaseModel):
    path: str = Field(..., min_length=1)  # full absolute path returned from candidate list


@app.post("/delete-file")
async def delete_file(body: DeleteFileRequest):
    """Delete a file by its full path (must be under home directory)."""
    target = Path(body.path).resolve()
    home = Path.home().resolve()

    # Safety: only allow deleting files inside the user's home directory
    if not str(target).startswith(str(home)):
        raise HTTPException(status_code=403, detail="Can only delete files inside your home directory")
    if not target.exists():
        raise HTTPException(status_code=404, detail=f"File not found: {target.name}")
    if not target.is_file():
        raise HTTPException(status_code=400, detail="Path is not a file")

    try:
        entry = ledger.snapshot_deletion(str(target), target.name)
        target.unlink()
        if entry:
            ledger.push(entry)
        logger.info("Deleted file: %s", target)
        return {"deleted": str(target), "filename": target.name}
    except Exception as e:
        logger.exception("Failed to delete file")
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/undo")
async def undo_last_action():
    """
    Directly reverse the last state-mutating action without going through the LLM.
    Called by the GUI Undo button and Cmd+Z keyboard shortcut.
    """
    success, message = ledger.undo()
    return {"success": success, "message": message}


class ScheduleMeetingRequest(BaseModel):
    title:            str
    attendees:        str = ""
    date:             str = ""
    time:             str = ""
    duration_minutes: int = 60
    location:         str = ""

class OpenAppRequest(BaseModel):
    app:    str
    action: str = ""

class OpenFileRequest(BaseModel):
    path: str


@app.post("/schedule-meeting")
async def schedule_meeting(body: ScheduleMeetingRequest):
    """Generate a .ics calendar file and open it in the system Calendar app."""
    import uuid
    import subprocess
    from datetime import datetime, timedelta

    try:
        dt_str = f"{body.date}T{body.time or '09:00'}:00"
        dt_start = datetime.fromisoformat(dt_str)
    except ValueError:
        dt_start = datetime.now().replace(hour=9, minute=0, second=0, microsecond=0)

    dt_end = dt_start + timedelta(minutes=body.duration_minutes or 60)
    fmt = "%Y%m%dT%H%M%S"

    attendee_lines = "\n".join(
        f"ATTENDEE;CN={a.strip()}:mailto:{a.strip()}"
        if "@" in a else f"ATTENDEE;CN={a.strip()}:mailto:unknown"
        for a in body.attendees.split(",") if a.strip()
    )

    ics = (
        "BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//Copilot//EN\n"
        "BEGIN:VEVENT\n"
        f"UID:{uuid.uuid4()}@copilot\n"
        f"DTSTAMP:{datetime.utcnow().strftime(fmt)}Z\n"
        f"DTSTART:{dt_start.strftime(fmt)}\n"
        f"DTEND:{dt_end.strftime(fmt)}\n"
        f"SUMMARY:{body.title}\n"
        + (f"LOCATION:{body.location}\n" if body.location else "")
        + (attendee_lines + "\n" if attendee_lines else "")
        + "END:VEVENT\nEND:VCALENDAR"
    )

    ics_path = Path(tempfile.gettempdir()) / "copilot_meeting.ics"
    entry = ledger.snapshot(str(ics_path), "SCHEDULE_MEETING", body.title)
    ics_path.write_text(ics, encoding="utf-8")
    ledger.push(entry)

    try:
        subprocess.run(["open", str(ics_path)], check=True, timeout=5)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not open Calendar: {e}")

    logger.info("Meeting scheduled: %s on %s", body.title, body.date)
    return {"ok": True, "title": body.title, "date": body.date, "time": body.time}


@app.post("/open-app")
async def open_app_endpoint(body: OpenAppRequest):
    """Open a macOS application by name, optionally performing a search action."""
    import subprocess
    import urllib.parse

    browser_apps = {"google chrome", "chrome", "safari", "firefox", "microsoft edge", "edge", "arc"}
    app_lower = body.app.lower()

    try:
        if body.action and any(b in app_lower for b in browser_apps):
            # Extract search terms and open a search URL
            search_terms = body.action.lower().replace("search for", "").replace("search", "").strip()
            url = f"https://www.bing.com/search?q={urllib.parse.quote(search_terms)}"
            subprocess.run(["open", "-a", body.app, url], check=True, timeout=5)
        else:
            subprocess.run(["open", "-a", body.app], check=True, timeout=5)
    except subprocess.CalledProcessError:
        raise HTTPException(status_code=404, detail=f"App not found: {body.app}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    logger.info("Opened app: %s (action: %s)", body.app, body.action)
    return {"ok": True, "app": body.app}


@app.post("/open-file")
async def open_file_endpoint(body: OpenFileRequest):
    """Open a file with its default system application."""
    import subprocess
    file_path = Path(body.path).expanduser()
    if not file_path.exists():
        raise HTTPException(status_code=404, detail=f"File not found: {body.path}")
    try:
        if sys.platform == "darwin":
            subprocess.run(["open", str(file_path)], check=True, timeout=5)
        elif sys.platform == "linux":
            subprocess.run(["xdg-open", str(file_path)], check=True, timeout=5)
        else:
            subprocess.run(["start", str(file_path)], shell=True, check=True, timeout=5)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"ok": True, "path": str(file_path)}


@app.post("/agent/run")
async def agent_run(body: AgentRequest):
    """Full Copilot agent pipeline streamed as Server-Sent Events."""
    conv_id = body.conversation_id
    history = [{"role": m.role, "content": m.content} for m in body.history]

    async def event_gen():
        assistant_content = ""
        try:
            async for chunk in run_agent_stream(body.query, history):
                # Intercept final events to capture the assistant's reply for storage
                if chunk.startswith("data: "):
                    try:
                        ev = _json.loads(chunk[6:].strip())
                        step = ev.get("step")
                        if step == "result":
                            assistant_content = str(ev.get("payload", ""))
                        elif step == "image":
                            assistant_content = f"[Generated image: \"{(ev.get('prompt') or '')[:100]}\". Saved and can be attached to emails.]"
                        elif step == "action_card":
                            pl = ev.get("payload", {})
                            assistant_content = f"[Proposed {ev.get('action')}: {_json.dumps(pl)}]"
                    except Exception:
                        pass
                yield chunk
        finally:
            # Persist messages to disk after stream completes
            if conv_id and assistant_content:
                _, was_first = chats_store.append_messages(conv_id, [
                    {"role": "user",      "content": body.query},
                    {"role": "assistant", "content": assistant_content},
                ])
                # Generate a smart title from the first user message
                if was_first:
                    title = await generate_chat_title(body.query)
                    chats_store.update_title(conv_id, title)

    return StreamingResponse(
        event_gen(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no", "Connection": "keep-alive"},
    )


# ── Chat management ───────────────────────────────────────────────────────────

@app.get("/chats")
async def list_chats():
    return chats_store.list_chats()

@app.post("/chats")
async def create_chat():
    return chats_store.create_chat()

@app.get("/chats/{chat_id}")
async def get_chat(chat_id: str):
    chat = chats_store.get_chat(chat_id)
    if not chat:
        raise HTTPException(status_code=404, detail="Chat not found")
    return chat

class RenameChatRequest(BaseModel):
    title: str

@app.patch("/chats/{chat_id}")
async def rename_chat(chat_id: str, body: RenameChatRequest):
    title = body.title.strip()
    if title:
        chats_store.update_title(chat_id, title)
    return {"ok": True}

@app.delete("/chats/{chat_id}")
async def delete_chat(chat_id: str):
    chats_store.delete_chat(chat_id)
    return {"ok": True}
